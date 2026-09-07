import io
import zipfile
from collections.abc import Callable
from pathlib import Path

from docx import Document


def build_docx(
    path: Path,
    *,
    paragraphs: list[str] | None = None,
    headings: list[str] | None = None,
    header_text: str | None = None,
    footer_text: str | None = None,
    image: bytes | None = None,
) -> Path:
    document = Document()
    for heading in headings or []:
        document.add_heading(heading, level=1)
    for paragraph in paragraphs or []:
        document.add_paragraph(paragraph)
    if image is not None:
        document.add_picture(io.BytesIO(image))
    section = document.sections[0]
    if header_text is not None:
        section.header.paragraphs[0].text = header_text
    if footer_text is not None:
        section.footer.paragraphs[0].text = footer_text
    document.save(path)
    return path


def rewrite_zip(path: Path, transform: Callable[[zipfile.ZipFile, zipfile.ZipInfo], bytes]) -> Path:
    with zipfile.ZipFile(path) as archive:
        entries = [
            (info, transform(archive, info)) for info in archive.infolist() if not info.is_dir()
        ]
    with zipfile.ZipFile(path, "w") as archive:
        for info, data in entries:
            info.file_size = len(data)
            archive.writestr(info, data)
    return path
